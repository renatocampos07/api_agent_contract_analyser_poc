import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.append(str(ROOT))

from app.analysis.docx_comments import _collect_run_spans  # noqa: E402
from docx import Document  # noqa: E402


DOC_PATH = ROOT / "data" / "processed" / "CONTRATO COM MARCA DE COMENTÁRIOS 1_validado.docx"

doc = Document(str(DOC_PATH))

with open(DOC_PATH, 'rb') as fh:
    import zipfile
    data = zipfile.ZipFile(fh).read('word/document.xml').decode('utf-8')


def comment_pos(cid: str) -> int:
    return data.find(f'w:commentRangeStart w:id="{cid}"')


print("Comment positions:")
ids = ["149", "150", "151", "152", "153", "158"]
for cid in ids:
    print(cid, comment_pos(cid))

print("\nPlaceholder positions:")
placeholders = ["XXXXXXX", "INDICAR ENDEREÇO", "00000000000000", "XX/XX/XXXX", "endereço da unidade CONTRATANTE"]
for ph in placeholders:
    print(ph, data.find(ph))


for idx, paragraph in enumerate(doc.paragraphs):
    if "XXXX" in paragraph.text or "Tel:" in paragraph.text or "endereço da unidade" in paragraph.text:
        print("\n== paragraph", idx, "==")
        print(paragraph.text)
        full_text, spans = _collect_run_spans(paragraph)
        print("full_text:", full_text)
        for run, start, end in spans:
            print(f"  span {start}-{end}: {repr(run.text)}")
